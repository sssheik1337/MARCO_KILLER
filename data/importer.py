"""Модуль импорта прайс-листов и остатков из XLSX/CSV в словари."""
from __future__ import annotations

import io
import logging
import re
import sqlite3
from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import BinaryIO, Union

import chardet
import pandas as pd
from openpyxl import load_workbook

from config import DB_PATH

SourceType = Union[str, BinaryIO]


class FormatError(Exception):
    """Исключение для ошибок формата входного файла."""


class SchemaMismatchError(Exception):
    """Исключение для несоответствия структуры таблицы ожидаемому формату."""

    def __init__(self, table_type: str, missing: list):
        super().__init__()
        self.table_type = table_type
        self.missing = missing
        logger.error(f"[SCHEMA ERROR] type={table_type} missing={self.missing}")


class ImportErrorFriendly(Exception):
    """Дружелюбное исключение для ошибок импорта (унифицированное)."""

    def __init__(
        self,
        title: str | None = None,
        details: str | None = None,
        template: str | None = None,
        reason: str | None = None,
        preview: list | None = None,
        total: int | None = None,
    ):
        super().__init__(title or reason)
        self.title = title
        self.details = details
        self.template = template
        self.reason = reason
        self.preview = preview or []
        self.total = total


@dataclass
class ParsedResult:
    """Результат разбора с валидными строками и предупреждениями."""

    items: list[dict]
    warnings: list[str]


TABLE_SCHEMAS = {
    "fabrics_catalog": {
        "required_columns": [
            "name",
            "country",
            "fabric_type",
            "segment",
            "wholesale_roll",
            "wholesale_piece",
        ],
        "template_path": "templates/fabrics_catalog_example.xlsx",
    },
    "hardware_catalog": {
        "required_columns": [
            "Артикул",
            "Наименование",
            "Коллекция",
            "Статус",
            "Кратность",
            "Бренд (Страна)",
            "Ед.",
            "Валюта",
            "РРЦ",
            "Оптовая",
        ],
        "template_path": "templates/hardware_catalog_example.xlsx",
    },
    "stock_fabrics_spb": {
        "required_columns": [
            "Номенклатура",
            "Остаток",
            "Свободный остаток",
        ],
        "template_path": "templates/stock_fabrics_spb_example.xlsx",
    },
    "stock_fabrics_msk": {
        "required_columns": [
            "Код товара",
            "Вид номенклатуры",
            "Тип номенклатуры",
            "Артикул",
            "Номенклатура",
            "Программа",
            "Наличие",
            "Ед.",
            "Доп. инфо.",
            "Дата прихода",
        ],
        "template_path": "templates/stock_fabrics_msk_example.xlsx",
    },
    "stock_hardware_spb": {
        "required_columns": [
            "Номенклатура",
            "Остаток",
            "Свободный остаток",
        ],
        "template_path": "templates/stock_hardware_spb_example.xlsx",
    },
    "stock_hardware_msk": {
        "required_columns": [
            "Код товара",
            "Вид номенклатуры",
            "Тип номенклатуры",
            "Артикул",
            "Номенклатура",
            "Программа",
            "Наличие",
            "Ед.",
            "Доп. инфо.",
            "Дата прихода",
            "Резерв",
        ],
        "template_path": "templates/stock_hardware_msk_example.xlsx",
    },
}


logger = logging.getLogger(__name__)


def classify_row_for_section(section: str, row: dict) -> bool:
    """Проверяет, соответствует ли строка разделу остатков (ткани/фурнитура)."""

    unit = str(row.get("Ед.") or row.get("Ед") or row.get("unit") or "").strip().lower()

    if section == "fabrics":
        allowed_units_fabrics = {
            "м",
            "м.",
            "метр",
            "метры",
            "m",
            "ед. хранения",
            "шт",
            "банк",
        }

        kind_text = str(row.get("Вид номенклатуры") or row.get("kind") or "").strip().lower()

        if "технические материалы" in kind_text:
            return True

        hardware_markers = {
            "фурнитура",
            "аксессуар",
            "светильник",
            "система",
            "петля",
            "направляющая",
        }

        if any(marker in kind_text for marker in hardware_markers) and unit not in allowed_units_fabrics:
            return False

        return True

    if section == "hardware":
        # Для фурнитуры пропускаем все непустые строки без дополнительной фильтрации
        has_value = any(str(value).strip() for value in row.values() if value is not None)
        return bool(has_value)

    return True


MAX_PREVIEW = 10


def _wrap_import_error(table_type: str, exc: Exception) -> ImportErrorFriendly:
    """Формирует дружелюбное исключение с учётом шаблона таблицы."""

    template = TABLE_SCHEMAS.get(table_type, {}).get("template_path")
    return ImportErrorFriendly(
        title=f"Не удалось импортировать таблицу {table_type}",
        details=str(exc),
        template=template,
    )


# ---------------------------------- служебные функции ----------------------------------

def _reset_stream(source: SourceType) -> SourceType:
    """Возвращает поток в начало, если он поддерживает seek."""

    if hasattr(source, "seek"):
        source.seek(0)
    return source


def _normalize_header(value: object) -> str:
    """Нормализует заголовок колонки: нижний регистр и одиночные пробелы."""

    return " ".join(str(value or "").strip().lower().split())


def _string(value: object) -> str | None:
    """Возвращает строку без лишних пробелов или None для пустых/NaN значений."""

    if isinstance(value, (list, tuple, set, pd.Series, pd.Index)):
        for item in value:
            result = _string(item)
            if result is not None:
                return result
        return None

    if value is None:
        return None

    try:
        if pd.isna(value):
            return None
    except TypeError:
        pass

    text = str(value).strip()
    return text or None


def _raw_text_value(value: object) -> str | None:
    """Возвращает строковое значение без очистки, если оно заполнено."""

    if value is None:
        return None

    try:
        if pd.isna(value):
            return None
    except TypeError:
        pass

    text = str(value)
    return text if text != "" else None


def _number(value: object) -> float | None:
    """Преобразует значение в float, очищая пробелы, запятые и валютные суффиксы."""

    if isinstance(value, (list, tuple, set, pd.Series, pd.Index)):
        for item in value:
            result = _number(item)
            if result is not None:
                return result
        return None

    if value is None:
        return None

    try:
        if pd.isna(value):
            return None
    except TypeError:
        pass

    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)

    text = str(value).strip()
    if not text:
        return None

    normalized = text.replace("\xa0", " ")
    cleaned = re.sub(r"[^0-9,\.\-]", "", normalized)
    if cleaned in {"", "-", ",", "."}:
        return None

    negative = cleaned.startswith("-")
    if negative:
        cleaned = cleaned[1:]

    cleaned = cleaned.strip()
    if not cleaned:
        return None

    last_comma = cleaned.rfind(",")
    last_dot = cleaned.rfind(".")

    number_text = cleaned
    decimal_sep = None
    thousands_sep = None

    if last_comma != -1 or last_dot != -1:
        if last_comma > last_dot:
            decimal_sep = ","
            thousands_sep = "." if last_dot != -1 else None
        elif last_dot > last_comma:
            decimal_sep = "."
            thousands_sep = "," if last_comma != -1 else None
        else:
            decimal_sep = "," if last_comma != -1 else "."

        if thousands_sep:
            number_text = number_text.replace(thousands_sep, "")
        if decimal_sep != ".":
            number_text = number_text.replace(decimal_sep, ".")

    number_text = number_text.replace(" ", "")

    try:
        number = float(number_text)
    except (TypeError, ValueError):
        return None

    if negative:
        number = -number

    return number


def _number_or_error(value: object, column_title: str) -> float | None:
    """Пытается преобразовать значение в число или сообщает об ошибке."""

    number = _number(value)
    if number is not None:
        return number

    text = _string(value)
    if text not in (None, "", "-", "—"):
        raise ImportErrorFriendly(
            title="Некорректное значение числа",
            details=f"Столбец «{column_title}» содержит некорректное значение.",
        )

    return None


def _detect_format(source: SourceType) -> str:
    """Определяет формат файла по расширению или сигнатуре."""

    name = getattr(source, "name", None) if not isinstance(source, str) else source
    prefix_bytes: bytes | None = None

    if isinstance(source, str):
        try:
            with open(source, "rb") as handle:
                prefix_bytes = handle.read(4)
        except OSError:
            prefix_bytes = None
    elif hasattr(source, "read"):
        current = _reset_stream(source)
        prefix = current.read(4)
        _reset_stream(current)
        prefix_bytes = prefix.encode() if isinstance(prefix, str) else prefix

    def _xls_not_supported() -> None:
        raise ImportErrorFriendly(
            title="Формат XLS не поддерживается",
            details="Используйте XLSX или CSV",
        )

    if name:
        lowered = name.lower()
        if lowered.endswith(".xls"):
            _xls_not_supported()
        if lowered.endswith(".xlsx"):
            if prefix_bytes is not None and not prefix_bytes.startswith(b"PK"):
                _xls_not_supported()
            return "xlsx"
        if lowered.endswith(".csv"):
            return "csv"

    if prefix_bytes is not None and prefix_bytes.startswith(b"PK"):
        return "xlsx"

    return "csv"


def _load_xlsx_rows(source: SourceType) -> list[list[object]]:
    """Читает XLSX с вычисленными формулами и возвращает список строк."""

    current = _reset_stream(source)
    try:
        workbook = load_workbook(current, data_only=True, read_only=True)
    except KeyError as e:
        raise ImportErrorFriendly(
            reason="bad_xlsx",
            preview=[str(e)],
            total=0,
        ) from e
    except Exception as exc:  # noqa: BLE001
        raise ImportErrorFriendly(
            reason="bad_xlsx",
            preview=[str(exc)],
            total=0,
        ) from exc

    try:
        sheet = workbook.active
        rows: list[list[object]] = []
        max_columns = sheet.max_column or 0
        for row in sheet.iter_rows(values_only=True):
            values = list(row)
            if max_columns and len(values) < max_columns:
                values.extend([None] * (max_columns - len(values)))
            rows.append(values)
    finally:
        workbook.close()
    return rows


def _load_csv_rows(source: SourceType) -> list[list[object]]:
    """Читает CSV без заголовков и возвращает список строк."""

    try:
        if isinstance(source, str):
            with open(source, "rb") as handle:
                data = handle.read()
        else:
            current = _reset_stream(source)
            data = current.read()
            _reset_stream(current)

        detection = chardet.detect(data)
        encoding = detection.get("encoding") or "utf-8"
        buffer = io.BytesIO(data)
        df = pd.read_csv(buffer, header=None, encoding=encoding)
        return df.where(pd.notna(df), None).values.tolist()
    except UnicodeDecodeError as exc:
        logger.error("Ошибка декодирования CSV", exc_info=True)
        raise ImportErrorFriendly(
            title="Ошибка чтения CSV",
            details="Некорректная кодировка файла.",
        ) from exc
    except pd.errors.ParserError as exc:  # type: ignore[attr-defined]
        logger.error("Ошибка парсинга CSV", exc_info=True)
        raise ImportErrorFriendly(
            title="Ошибка чтения CSV",
            details="Некорректная структура CSV-файла.",
        ) from exc
    except Exception as exc:  # noqa: BLE001
        logger.error("Неизвестная ошибка при чтении CSV", exc_info=True)
        raise ImportErrorFriendly(
            title="Ошибка чтения CSV-файла",
            details="Проверьте корректность данных.",
        ) from exc


def _load_rows(source: SourceType) -> list[list[object]]:
    """Возвращает все строки файла вне зависимости от формата."""

    fmt = _detect_format(source)
    if fmt == "xlsx":
        return _load_xlsx_rows(source)
    if fmt == "csv":
        return _load_csv_rows(source)
    raise FormatError("Формат файла не поддерживается. Используйте XLSX или CSV")


def _ensure_columns(columns: dict[str, object], required: list[str]) -> None:
    """Проверяет наличие обязательных колонок, выбрасывая осмысленную ошибку."""

    for title in required:
        if _normalize_header(title) not in columns:
            raise ImportErrorFriendly(
                title="Загруженная таблица не соответствует формату",
                details=f"Отсутствует столбец «{title}».",
            )


def _validate_headers(table_type: str, headers: list[str]) -> None:
    """Проверяет структуру таблицы по ожидаемой схеме."""

    schema = TABLE_SCHEMAS[table_type]
    required = schema["required_columns"]
    if not headers or all((h or "").strip() == "" for h in headers):
        missing = required
    else:
        missing = [c for c in required if c not in headers]

    if missing:
        logger.error(
            f"Ошибка структуры ({table_type}): отсутствуют колонки: {missing}",
        )
        raise ImportErrorFriendly(
            reason="missing_columns",
            preview=missing,
            template=TABLE_SCHEMAS[table_type]["template_path"],
        )


def _row_value(row: list[object], index: int) -> object:
    """Безопасно получает значение ячейки по индексу."""

    return row[index] if 0 <= index < len(row) else None


def _load_catalog_index(table: str) -> tuple[dict[str, int], dict[str, int]]:
    """Возвращает словари поиска по артикулу и названию для каталога."""

    article_map: dict[str, int] = {}
    name_map: dict[str, int] = {}
    with sqlite3.connect(DB_PATH) as conn:
        try:
            cur = conn.execute(f"SELECT id, article, name FROM {table}")
            rows = cur.fetchall()
            for catalog_id, article, name in rows:
                if article:
                    article_map[_normalize_header(article)] = int(catalog_id)
                if name:
                    name_map[_normalize_header(name)] = int(catalog_id)
            return article_map, name_map
        except sqlite3.OperationalError:
            return article_map, name_map


def _find_catalog_id(
    article: str | None, name: str | None, article_map: dict[str, int], name_map: dict[str, int]
) -> int | None:
    """Ищет id каталога по артикулу или названию."""

    if article:
        normalized = _normalize_header(article)
        if normalized in article_map:
            return article_map[normalized]
    if name:
        normalized = _normalize_header(name)
        if normalized in name_map:
            return name_map[normalized]
    return None


# ---------------------------------- парсинг каталогов ----------------------------------

def parse_fabrics_catalog(stream: SourceType) -> list[dict]:
    """Парсит каталог тканей с динамическими диапазонами цен."""

    try:
        table_type = "fabrics_catalog"
        rows = _load_rows(stream)
        if not rows:
            return []

        logger.info(f"Импорт {table_type} для города -")
        logger.debug("Всего сырых строк в файле: %s", len(rows))

        header_index = None
        for idx, row in enumerate(rows):
            first_cell = _normalize_header(_row_value(row, 0))
            if first_cell == "наименование коллекции":
                header_index = idx
                break

        if header_index is None:
            raise ImportErrorFriendly(
                title="В таблице отсутствует строка заголовков",
                details="Не найдена колонка 'Наименование коллекции'.",
                template=TABLE_SCHEMAS[table_type]["template_path"],
            )

        header_row = rows[header_index]
        top_row = rows[header_index - 1] if header_index - 1 >= 0 else None

        base_positions: dict[str, int] = {
            "name": 0,
            "country": 1,
            "fabric_type": 2,
            "segment": 3,
            "wholesale_roll": 4,
            "wholesale_piece": 5,
        }

        # Начинаем разбор диапазонов с колонки 6, как только над шапкой указано имя диапазона.
        def _normalize_range_name(value: str) -> str:
            cleaned = re.sub(r"[^0-9]+", "_", value)
            cleaned = cleaned.strip("_")
            return cleaned

        def _is_role_cell(value: object, expected: set[str]) -> bool:
            """Проверяет, соответствует ли ячейка тексту "ролик"/"отрез"."""

            normalized = _normalize_header(value)
            return normalized in expected

        def _is_range_title(value: object) -> bool:
            """Определяет, подходит ли верхняя ячейка под шаблон диапазона."""

            text = _string(value)
            if not text:
                return False
            return bool(re.search(r"\d+\s*[-–]\s*\d+", text))

        range_columns: dict[str, dict[str, int]] = {}
        col = 6
        max_len = max(len(header_row), len(top_row or []))

        while col < max_len:
            range_title_cell = _row_value(top_row, col) if top_row else None
            if range_title_cell is None or _string(range_title_cell) in (None, ""):
                break

            if not _is_range_title(range_title_cell):
                raise ImportErrorFriendly(
                    title="Некорректное название диапазона",
                    details=(
                        "Ячейка над колонкой не содержит диапазон ширины (пример: 85-90)."
                    ),
                    template=TABLE_SCHEMAS[table_type]["template_path"],
                )

            range_title = _string(range_title_cell) or ""
            range_key = _normalize_range_name(range_title)
            if not range_key:
                raise ImportErrorFriendly(
                    title="Некорректное название диапазона",
                    details="Ячейка диапазона не содержит числового интервала (пример: 85-90).",
                    template=TABLE_SCHEMAS[table_type]["template_path"],
                )

            roll_idx = col
            piece_idx = col + 1

            roll_ok = _is_role_cell(_row_value(header_row, roll_idx), {"ролик", "ролл", "roll"})
            piece_ok = _is_role_cell(
                _row_value(header_row, piece_idx), {"отрез", "отр", "piece"}
            )

            if not (roll_ok and piece_ok):
                raise ImportErrorFriendly(
                    title="Некорректная структура диапазона",
                    details=(
                        "Ожидаются подписи 'РОЛИК' и 'отрез' под диапазоном "
                        f"{range_title}, но они не найдены."
                    ),
                    template=TABLE_SCHEMAS[table_type]["template_path"],
                )

            range_columns[range_key] = {"roll": roll_idx, "piece": piece_idx}
            col += 2

        if not range_columns:
            raise ImportErrorFriendly(
                title="Не найдены диапазоны цен",
                details="После обязательных колонок отсутствуют диапазоны ширины ткани.",
                template=TABLE_SCHEMAS[table_type]["template_path"],
            )

        data_start = header_index + 1

        items: list[dict] = []
        for row in rows[data_start:]:
            name = _string(_row_value(row, base_positions["name"]))
            if not name:
                logger.debug(
                    "Строка пропущена: отсутствует 'Наименование коллекции' в %s",
                    row,
                )
                continue

            special = None
            status_search_start = col - 1
            last_index = len(row) - 1
            for idx in range(last_index, status_search_start - 1, -1):
                cell = _row_value(row, idx)
                if cell is None:
                    continue
                if isinstance(cell, str):
                    candidate = cell.strip()
                    if not candidate:
                        continue
                    if candidate.startswith("="):
                        continue
                    if _number(candidate) is not None:
                        continue
                    if not any(ch.isalpha() for ch in candidate):
                        continue
                    special = candidate
                    break
                if isinstance(cell, (int, float)):
                    continue

            item = {
                "city": "all",
                "section": "fabrics",
                "category": _string(_row_value(row, base_positions["segment"])),
                "subcategory": None,
                "name": name,
                "country": _string(_row_value(row, base_positions["country"])),
                "fabric_type": _string(
                    _row_value(row, base_positions["fabric_type"])
                ),
                "segment": _string(_row_value(row, base_positions["segment"])),
                "wholesale_roll": _raw_text_value(
                    _row_value(row, base_positions["wholesale_roll"])
                ),
                "wholesale_piece": _raw_text_value(
                    _row_value(row, base_positions["wholesale_piece"])
                ),
                "special": special,
                "image_url": None,
            }

            for range_key, cols in range_columns.items():
                roll_value = _row_value(row, cols["roll"])
                piece_value = _row_value(row, cols["piece"])

                field_name_roll = f"price_roll_{range_key}"
                field_name_piece = f"price_piece_{range_key}"

                item[field_name_roll] = _number_or_error(roll_value, field_name_roll)
                item[field_name_piece] = _number_or_error(
                    piece_value, field_name_piece
                )

            items.append(item)
            logger.debug("Строка добавлена: %s", item)

        logger.info(f"Получено валидных записей: {len(items)}")
        logger.debug("Итоговое количество валидных строк каталога: %s", len(items))

        return items
    except ImportErrorFriendly:
        raise
    except Exception as exc:  # noqa: BLE001
        logger.error("Ошибка импорта каталога тканей", exc_info=True)
        raise _wrap_import_error(table_type, exc) from exc


def parse_hardware_catalog(stream: SourceType) -> list[dict]:
    """Парсит каталог фурнитуры строго по заполненности столбцов."""

    try:
        table_type = "hardware_catalog"
        rows = _load_rows(stream)
        if not rows:
            return []

        logger.info("Импорт %s для города all", table_type)

        working_rows = rows[10:] if len(rows) > 10 else rows
        if not working_rows:
            raise ImportErrorFriendly(
                reason="missing_columns",
                details="Не найдена строка заголовков каталога фурнитуры.",
            )

        header_row_full: list[object] | None = None
        header_cells: list[object] = []
        normalized_headers: dict[str, int] = {}
        for row in working_rows:
            candidate_cells = row[1:]
            normalized = {_normalize_header(val): idx for idx, val in enumerate(candidate_cells)}
            required_titles = [
                "Артикул",
                "Наименование",
                "Коллекция",
                "Статус",
                "Кратность",
                "Бренд (Страна)",
                "Ед.",
                "Валюта",
                "РРЦ",
                "Оптовая",
            ]
            try:
                _ensure_columns(normalized, required_titles)
            except ImportErrorFriendly:
                continue
            header_row_full = row
            header_cells = candidate_cells
            normalized_headers = normalized
            break

        if not header_cells:
            raise ImportErrorFriendly(
                reason="missing_columns",
                details="Не найдена строка заголовков каталога фурнитуры.",
            )

        headers = [(_string(val) or "").strip() for val in header_cells]
        _validate_headers(table_type, headers)

        header_index = working_rows.index(header_row_full)
        data_rows = working_rows[header_index + 1 :]
        trimmed_rows = [row[1 : len(header_cells) + 1] for row in data_rows]

        col_article = normalized_headers[_normalize_header("Артикул")]
        col_name = normalized_headers[_normalize_header("Наименование")]
        col_status = normalized_headers[_normalize_header("Статус")]
        col_multiplicity = normalized_headers[_normalize_header("Кратность")]
        col_brand = normalized_headers[_normalize_header("Бренд (Страна)")]
        col_unit = normalized_headers[_normalize_header("Ед.")]
        col_currency = normalized_headers[_normalize_header("Валюта")]
        col_rrc = normalized_headers[_normalize_header("РРЦ")]
        col_opt = normalized_headers[_normalize_header("Оптовая")]

        current_cat1: str | None = None
        current_cat2: str | None = None
        current_cat3: str | None = None
        pending_cat2_list: list[str] = []
        goods_started = False
        items: list[dict] = []

        def _is_up_marker(value: object) -> bool:
            text = _string(value)
            return bool(text) and text.strip().lower() == "вверх"

        def _is_empty(value: object) -> bool:
            """Пустая ячейка или служебная надпись «вверх».

            «ВВЕРХ» игнорируется на уровне ячейки, чтобы не ломать разметку
            категорий, если ссылка соседствует с названием уровня.
            """

            text = _string(value)
            return text is None or text == "" or _is_up_marker(text)

        def _only_column_has_text(row_values: list[object], idx: int) -> bool:
            return bool(_string(row_values[idx])) and all(
                _is_empty(val) for pos, val in enumerate(row_values) if pos != idx
            )

        def _has_letters(value: object) -> bool:
            return isinstance(value, str) and any(ch.isalpha() for ch in value)

        category_totals: dict[str, dict[str | None, dict[str | None, int]]]
        category_totals = defaultdict(lambda: defaultdict(lambda: defaultdict(int)))

        for row in trimmed_rows:
            if all(_is_empty(cell) for cell in row):
                continue

            article_cell = row[col_article]
            name_cell = row[col_name]
            status_cell = row[col_status]
            price_rrc_raw = row[col_rrc]
            price_opt_raw = row[col_opt]

            if (
                str(article_cell).strip().lower() == "артикул"
                and str(name_cell).strip().lower() == "наименование"
            ):
                continue

            price_rrc = _number(price_rrc_raw)
            price_opt = _number(price_opt_raw)
            has_price = not (_is_empty(price_rrc_raw) and _is_empty(price_opt_raw))

            is_cat_level1 = (
                _only_column_has_text(row, 0)
                and _string(row[0])
                and all(_is_empty(cell) for cell in row[1:])
            )
            is_cat_level2 = _only_column_has_text(row, 1) and all(
                _is_empty(cell) for cell in row[2:]
            )
            is_cat_level3 = _only_column_has_text(row, 2) and all(
                _is_empty(cell) for cell in row[3:]
            )

            if is_cat_level1:
                title = _string(row[0])
                if title and title.strip().lower() != "вверх":
                    current_cat1 = title
                    logger.info("[IMPORT] category_1: %s", current_cat1)
                else:
                    current_cat1 = None
                current_cat2 = None
                current_cat3 = None
                pending_cat2_list = []
                goods_started = False
                continue

            if is_cat_level2:
                title = _string(row[1])
                if title and title.strip().lower() != "вверх":
                    current_cat2 = title
                    logger.info("[IMPORT]   category_2: %s", current_cat2)
                    if not goods_started and current_cat2 not in pending_cat2_list:
                        pending_cat2_list.append(current_cat2)
                else:
                    current_cat2 = None
                current_cat3 = None
                continue

            if is_cat_level3:
                title = _string(row[2])
                if title and title.strip().lower() != "вверх":
                    if current_cat2 is None and not goods_started and pending_cat2_list:
                        current_cat2 = pending_cat2_list[0]
                        logger.info("[IMPORT]   category_2: %s", current_cat2)
                        logger.info(
                            "[IMPORT][FIX] implicit category_2 applied: %s", current_cat2
                        )
                    if title != current_cat3:
                        logger.info("[IMPORT]     category_3: %s", title)
                    current_cat3 = title
                else:
                    current_cat3 = None
                continue

            article = _string(article_cell)
            name = _string(name_cell)

            if not article or not name or not has_price or not current_cat1:
                continue

            status_text = _string(status_cell)
            if not _has_letters(status_text):
                status_text = None
            else:
                status_text = status_text.strip()

            item = {
                "city": "all",
                "section": "hardware",
                "category": current_cat1,
                "subcategory": current_cat2,
                "group": current_cat3,
                "article": article,
                "name": name,
                "special": status_text,
                "multiplicity": _string(row[col_multiplicity]),
                "brand_country": _string(row[col_brand]),
                "unit": _string(row[col_unit]),
                "currency": _string(row[col_currency]),
                "price_rrc": price_rrc,
                "price_opt": price_opt,
                "image_url": None,
            }

            items.append(item)
            goods_started = True
            category_totals[current_cat1][current_cat2][current_cat3] += 1

        for cat1, sub_map in category_totals.items():
            logger.info("[IMPORT] category_1: %s", cat1)
            for cat2, grp_map in sub_map.items():
                if cat2:
                    logger.info("[IMPORT]   category_2: %s", cat2)
                for cat3, count in grp_map.items():
                    if cat3:
                        logger.info("[IMPORT]     category_3: %s", cat3)
                    logger.info("[IMPORT]       products: %s", count)

        logger.info("Получено валидных записей каталога фурнитуры: %s", len(items))

        return items
    except ImportErrorFriendly:
        raise
    except Exception as exc:  # noqa: BLE001
        logger.error("Ошибка импорта каталога фурнитуры", exc_info=True)
        raise _wrap_import_error(table_type, exc) from exc


# ---------------------------------- парсинг остатков ----------------------------------

def parse_hardware_stock_msk(stream: SourceType) -> list[dict]:
    """Парсит остатки фурнитуры для Москвы."""

    try:
        table_type = "stock_hardware_msk"
        rows = _load_rows(stream)
        if not rows:
            return []

        city = "msk"
        logger.info("Импорт %s для города %s", table_type, city)

        required_titles = [
            "Код товара",
            "Вид номенклатуры",
            "Тип номенклатуры",
            "Артикул",
            "Номенклатура",
            "Программа",
            "Наличие",
            "Ед.",
            "Доп. инфо.",
            "Дата прихода",
        ]
        optional_reserved = "Резерв"

        header_index = None
        columns: dict[str, str] = {}
        header_row: list[object] | None = None
        for idx, row in enumerate(rows):
            temp_columns = {_normalize_header(val): val for val in row}
            try:
                _ensure_columns(temp_columns, required_titles)
                header_index = idx
                columns = temp_columns
                header_row = row
                break
            except Exception:
                continue

        if header_index is None or header_row is None:
            temp_columns = {_normalize_header(val): val for val in rows[0]}
            _ensure_columns(temp_columns, required_titles)
            header_index = 0
            header_row = rows[0]
            columns = temp_columns

        headers = [(_string(val) or "").strip() for val in header_row]
        _validate_headers(table_type, headers)

        data_rows = rows[header_index + 1 :]
        df = pd.DataFrame(data_rows, columns=header_row)
        df = df.dropna(how="all")

        article_map, name_map = _load_catalog_index("hardware_catalog")
        reserved_idx = columns.get(_normalize_header(optional_reserved))

        records = df.to_dict(orient="records")
        foreign_rows: list[dict] = []
        items: list[dict] = []
        for record in records:
            row = dict(record)
            if not classify_row_for_section("hardware", row):
                foreign_rows.append(row)
                continue

            name = _string(row.get(columns["номенклатура"]))
            article = _string(row.get(columns["артикул"]))
            quantity = _number_or_error(row.get(columns["наличие"]), "Наличие")
            unit = _string(row.get(columns["ед."]))
            additional_info = _string(row.get(columns["доп. инфо."]))
            arrival_date = _string(row.get(columns["дата прихода"]))
            kind = _string(row.get(columns["вид номенклатуры"]))
            item_type = _string(row.get(columns["тип номенклатуры"]))
            code = _string(row.get(columns["код товара"]))

            if not name and not article:
                logger.warning("Запись пропущена: нет артикула и наименования")
                continue

            catalog_id = _find_catalog_id(article, name, article_map, name_map)

            item = {
                "city": "msk",
                "section": "hardware",
                "kind": kind,
                "item_type": item_type,
                "code": code,
                "catalog_id": catalog_id,
                "name": name,
                "article": article,
                "quantity": quantity,
                "free_quantity": None,
                "unit": unit,
                "arrival_date": arrival_date,
                "reserved": _string(row.get(reserved_idx)) if reserved_idx is not None else None,
                "extra_info": additional_info,
            }
            items.append(item)

        if foreign_rows:
            ratio = len(foreign_rows) / max(len(records), 1)
            preview = [_string(r.get("Номенклатура")) for r in foreign_rows[:MAX_PREVIEW]]
            if len(foreign_rows) <= 20 and ratio < 0.10:
                return ParsedResult(items=items, warnings=preview)

            raise ImportErrorFriendly(
                reason="wrong_section", preview=preview, total=len(foreign_rows)
            )

        logger.info("[IMPORT DONE] Type=%s City=%s Items=%s", table_type, city, len(items))

        return items
    except ImportErrorFriendly:
        raise
    except Exception as exc:  # noqa: BLE001
        logger.error("Ошибка импорта остатков фурнитуры (Москва)", exc_info=True)
        raise _wrap_import_error(table_type, exc) from exc


def parse_hardware_stock_spb(stream: SourceType) -> dict:
    """Сохраняет файл остатков фурнитуры СПБ без разбора содержимого."""

    target_dir = Path("data/stocks/spb/hardware")
    target_dir.mkdir(parents=True, exist_ok=True)
    file_path = target_dir / "hardware_stock_spb.xlsx"

    try:
        # Удаляем предыдущий файл, если он был сохранён ранее
        if file_path.exists():
            file_path.unlink()

        if hasattr(stream, "seek"):
            stream.seek(0)
        with open(file_path, "wb") as f:
            f.write(stream.read())

        logger.info(
            "Файл остатков фурнитуры СПБ сохранён без разбора: %s", file_path
        )
        return {"items": [], "imported": 0, "saved_path": str(file_path)}
    except Exception as exc:  # noqa: BLE001
        logger.error("Не удалось сохранить файл остатков фурнитуры СПБ", exc_info=True)
        raise ImportErrorFriendly(
            title="Не удалось сохранить файл остатков фурнитуры СПБ",
            details=str(exc),
            template=None,
        ) from exc


def parse_fabrics_stock_msk(stream: SourceType) -> dict:
    """Парсит остатки тканей для Москвы с мягкой фильтрацией посторонних строк."""

    try:
        table_type = "stock_fabrics_msk"
        rows = _load_rows(stream)
        if not rows:
            return []

        city = "msk"
        logger.info("Импорт %s для города %s", table_type, city)

        required_titles = [
            "Код товара",
            "Вид номенклатуры",
            "Тип номенклатуры",
            "Артикул",
            "Номенклатура",
            "Программа",
            "Наличие",
            "Ед.",
            "Доп. инфо.",
            "Дата прихода",
        ]

        header_index = None
        columns: dict[str, str] = {}
        header_row: list[object] | None = None
        for idx, row in enumerate(rows):
            temp_columns = {_normalize_header(val): val for val in row}
            try:
                _ensure_columns(temp_columns, required_titles)
                header_index = idx
                columns = temp_columns
                header_row = row
                break
            except Exception:
                continue

        if header_index is None or header_row is None:
            temp_columns = {_normalize_header(val): val for val in rows[0]}
            _ensure_columns(temp_columns, required_titles)
            header_index = 0
            header_row = rows[0]
            columns = temp_columns

        headers = [(_string(val) or "").strip() for val in header_row]
        _validate_headers(table_type, headers)

        data_rows = rows[header_index + 1 :]
        df = pd.DataFrame(data_rows, columns=header_row)
        df = df.dropna(how="all")

        records = df.to_dict(orient="records")
        foreign_rows: list[dict] = []
        items: list[dict] = []
        for record in records:
            row = dict(record)
            if not classify_row_for_section("fabrics", row):
                foreign_rows.append(row)
                continue

            name = _string(row.get(columns["номенклатура"]))
            article = _string(row.get(columns["артикул"]))
            code = _string(row.get(columns["код товара"]))
            if not name and not article:
                logger.warning("Запись пропущена: нет артикула и наименования")
                continue

            quantity = _number_or_error(row.get(columns["наличие"]), "Наличие")
            unit = _string(row.get(columns["ед."]))
            extra_info = _string(row.get(columns["доп. инфо."]))
            kind = _string(row.get(columns["вид номенклатуры"]))
            item_type = _string(row.get(columns["тип номенклатуры"]))

            if quantity is None:
                logger.warning("Запись пропущена: не указано количество")
                continue

            item = {
                "city": "msk",
                "section": "fabrics",
                "kind": kind,
                "item_type": item_type,
                "code": code,
                "article": article,
                "name": name,
                "quantity": quantity,
                "unit": unit,
                "extra_info": extra_info,
            }
            items.append(item)

        if foreign_rows:
            ratio = len(foreign_rows) / max(len(records), 1)
            preview = [_string(r.get("Номенклатура")) for r in foreign_rows[:MAX_PREVIEW]]
            if len(foreign_rows) <= 20 and ratio < 0.10:
                return ParsedResult(items=items, warnings=preview)

            raise ImportErrorFriendly(
                reason="wrong_section", preview=preview, total=len(foreign_rows)
            )

        logger.info("[IMPORT DONE] Type=%s City=%s Items=%s", table_type, city, len(items))

        return {"items": items, "imported": len(items), "skipped": []}
    except ImportErrorFriendly:
        raise
    except Exception as exc:  # noqa: BLE001
        logger.error("Ошибка импорта остатков тканей (Москва)", exc_info=True)
        raise _wrap_import_error(table_type, exc) from exc


def parse_fabrics_stock_spb(stream: SourceType) -> dict:
    """Парсит остатки тканей для Санкт-Петербурга без фильтрации строк."""

    try:
        table_type = "stock_fabrics_spb"
        rows = _load_rows(stream)
        if not rows:
            return []

        city = "spb"
        logger.info("Импорт %s для города %s", table_type, city)

        base_headers = ["Номенклатура", "Остаток", "Свободный остаток"]

        header_index = None
        main_header: list[object] | None = None
        sub_header: list[object] | None = None
        for idx, row in enumerate(rows):
            normalized = [_normalize_header(val) for val in row]
            if all(_normalize_header(title) in normalized for title in base_headers):
                header_index = idx
                main_header = row
                if idx + 1 < len(rows):
                    sub_header = rows[idx + 1]
                break

        if header_index is None or main_header is None:
            raise ImportErrorFriendly(
                title="Загруженная таблица не соответствует формату",
                details=(
                    "Структура таблицы остатков тканей СПБ не распознана.\n"
                    "Требуемые колонки: Номенклатура, Остаток, Свободный остаток."
                ),
                template=TABLE_SCHEMAS[table_type]["template_path"],
            )

        max_len = max(len(main_header), len(sub_header or []))
        combined_headers: list[object] = []
        for i in range(max_len):
            top = main_header[i] if i < len(main_header) else None
            bottom = sub_header[i] if sub_header and i < len(sub_header) else None
            top_str = _string(top)
            bottom_str = _string(bottom)

            if bottom_str:
                combined = f"{top_str or ''} ({bottom_str})".strip()
            else:
                combined = top_str or None
            combined_headers.append(combined)

        normalized_headers = {
            _normalize_header(val): val for val in combined_headers if val is not None
        }

        headers: list[str] = []
        for value in combined_headers:
            normalized = _normalize_header(value)
            if normalized == _normalize_header("Номенклатура"):
                headers.append("Номенклатура")
            elif normalized in {
                _normalize_header("Остаток (В ед. хранения)"),
                _normalize_header("Остаток"),
            }:
                headers.append("Остаток")
            elif normalized in {
                _normalize_header("Свободный остаток (В ед. хранения)"),
                _normalize_header("Свободный остаток"),
            }:
                headers.append("Свободный остаток")

        _validate_headers(table_type, headers)

        data_rows = rows[(header_index + 2 if sub_header else header_index + 1) :]
        df = pd.DataFrame(data_rows, columns=combined_headers)
        df = df.dropna(how="all")

        items: list[dict] = []
        name_key = normalized_headers[_normalize_header("Номенклатура")]
        quantity_key = normalized_headers.get(
            _normalize_header("Остаток (В ед. хранения)")
        ) or normalized_headers[_normalize_header("Остаток")]
        free_key = normalized_headers.get(
            _normalize_header("Свободный остаток (В ед. хранения)")
        ) or normalized_headers[_normalize_header("Свободный остаток")]

        for record in df.to_dict(orient="records"):
            name = record.get(name_key)
            quantity = record.get(quantity_key)
            free_quantity = record.get(free_key)

            items.append(
                {
                    "city": city,
                    "section": "fabrics",
                    "name": name if name is not None else "",
                    "code": None,
                    "quantity": quantity,
                    "free_quantity": free_quantity,
                    "unit": "м",
                    "kind": None,
                    "item_type": None,
                    "article": None,
                    "extra_info": None,
                }
            )

        logger.info("[IMPORT DONE] Type=%s City=%s Items=%s", table_type, city, len(items))

        return {"items": items, "imported": len(items), "skipped": []}
    except ImportErrorFriendly:
        raise
    except Exception as exc:  # noqa: BLE001
        logger.error("Ошибка импорта остатков тканей (СПБ)", exc_info=True)
        raise _wrap_import_error(table_type, exc) from exc


def parse_fabrics(stream: SourceType, city: str | None = None, section: str | None = None):
    """Совместимость со старым API: возвращает каталог тканей."""

    return parse_fabrics_catalog(stream)


def parse_hardware(stream: SourceType, city: str | None = None, section: str | None = None):
    """Совместимость со старым API: возвращает каталог фурнитуры."""

    return parse_hardware_catalog(stream)


def parse_stock(stream: SourceType, city: str | None = None, section: str | None = None):
    """Заглушка для остатков: требует явного указания раздела."""

    if section is None:
        raise ImportErrorFriendly(
            title="Не указан раздел остатков",
            details="Укажите раздел остатков (fabrics/hardware) для импорта.",
        )

    normalized = (section or "").lower()
    if normalized == "fabrics":
        if city == "msk":
            return parse_fabrics_stock_msk(stream)
        return parse_fabrics_stock_spb(stream)

    if normalized == "hardware":
        if city == "msk":
            return parse_hardware_stock_msk(stream)
        return parse_hardware_stock_spb(stream)

    raise ImportErrorFriendly(
        title="Неизвестный раздел остатков",
        details="Ожидается fabrics или hardware.",
    )


__all__ = [
    "parse_fabrics_catalog",
    "parse_hardware_catalog",
    "parse_fabrics_stock_spb",
    "parse_fabrics_stock_msk",
    "parse_hardware_stock_spb",
    "parse_hardware_stock_msk",
    "parse_fabrics",
    "parse_hardware",
    "parse_stock",
    "_string",
    "_number",
    "_number_or_error",
    "ImportErrorFriendly",
    "TABLE_SCHEMAS",
]
